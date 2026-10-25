import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl

# Paths
input_docx = r"C:\Users\mohammed saad\.gemini\antigravity\brain\5ea6a609-404f-416f-b8fc-79047196e534\.user_uploaded\media_1792931729858_71b15c53.docx"
output_project_docx = r"C:\Users\mohammed saad\OneDrive\Desktop\quiz-platform-auth\QUIZNOVA_AI_BLACK_BOOK_PROJECT_REPORT.docx"
output_desktop_docx = r"C:\Users\mohammed saad\OneDrive\Desktop\QUIZNOVA_AI_BLACK_BOOK_PROJECT_REPORT.docx"
diagrams_dir = r"C:\Users\mohammed saad\OneDrive\Desktop\quiz-platform-auth\Project_Diagrams_Word"

print("Loading original document...")
doc = docx.Document(input_docx)
body = doc._body._element

# Find Chapter 3 and Chapter 4 indices
ch3_start = -1
ch4_start = -1
ch4_element = None

for idx, child in enumerate(body):
    if isinstance(child, CT_P):
        p = docx.text.paragraph.Paragraph(child, doc)
        txt = p.text.strip()
        if 'CHAPTER 3: SYSTEM DESIGN' in txt and ch3_start == -1:
            ch3_start = idx
        elif 'CHAPTER 4 : UI SCREENS' in txt and ch4_start == -1:
            ch4_start = idx
            ch4_element = child

print(f"Chapter 3 starts at element: {ch3_start}, Chapter 4 starts at element: {ch4_start}")

# Remove elements between ch3_start and ch4_start
elements_to_remove = []
for idx in range(ch3_start + 1, ch4_start):
    elements_to_remove.append(body[idx])

for el in elements_to_remove:
    body.remove(el)

print(f"Removed {len(elements_to_remove)} legacy elements from Chapter 3.")

# Helper to style a table nicely
def style_table(table, header_bg="2B3A4A", zebra_bg="F4F6F9"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(table.rows):
        # Set cantSplit
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(r'<w:cantSplit xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))
        
        # Header repeat
        if r_idx == 0:
            trPr.append(parse_xml(r'<w:tblHeader xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))
            
        for c_idx, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            # Padding
            tcMar = parse_xml(r'<w:tcMar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                              r'<w:top w:w="120" w:type="dxa"/>'
                              r'<w:bottom w:w="120" w:type="dxa"/>'
                              r'<w:left w:w="160" w:type="dxa"/>'
                              r'<w:right w:w="160" w:type="dxa"/>'
                              r'</w:tcMar>')
            tcPr.append(tcMar)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            # Shading
            if r_idx == 0:
                shd = parse_xml(f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="{header_bg}"/>')
                tcPr.append(shd)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9.5)
                        run.font.name = "Arial"
            else:
                if r_idx % 2 == 0 and zebra_bg:
                    shd = parse_xml(f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="{zebra_bg}"/>')
                    tcPr.append(shd)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
                        run.font.name = "Arial"

# Helper builders that insert elements right before ch4_element
def add_p(text, bold=False, italic=False, size_pt=11, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=4, space_after=6, color_rgb=(0,0,0)):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.bold = bold
    run.font.italic = italic
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RGBColor(*color_rgb)
    ch4_element.addprevious(p._p)
    return p

def add_heading2(title):
    return add_p(title, bold=True, size_pt=14, space_before=16, space_after=6, color_rgb=(15, 23, 42))

def add_heading3(title):
    return add_p(title, bold=True, size_pt=12, space_before=12, space_after=4, color_rgb=(30, 41, 59))

def add_bullet(text, prefix="• "):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r_bullet = p.add_run(prefix)
    r_bullet.font.name = "Arial"
    r_bullet.font.bold = True
    r_bullet.font.size = Pt(10)
    r_text = p.add_run(text)
    r_text.font.name = "Arial"
    r_text.font.size = Pt(10)
    ch4_element.addprevious(p._p)
    return p

def add_diagram(filename, caption, width_in=6.0):
    img_path = os.path.join(diagrams_dir, filename)
    if not os.path.exists(img_path):
        print(f"Warning: Image {img_path} not found!")
        return None
        
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(10)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(width_in))
    ch4_element.addprevious(p_img._p)
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(16)
    run_cap = p_cap.add_run(caption)
    run_cap.font.name = "Arial"
    run_cap.font.italic = True
    run_cap.font.bold = True
    run_cap.font.size = Pt(9.5)
    run_cap.font.color.rgb = RGBColor(71, 85, 105)
    ch4_element.addprevious(p_cap._p)
    return p_img

def add_tbl(headers, rows, col_widths=None):
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    for c_idx, h in enumerate(headers):
        tbl.cell(0, c_idx).text = h
    for r_idx, r in enumerate(rows):
        for c_idx, val in enumerate(r):
            tbl.cell(r_idx + 1, c_idx).text = str(val)
    style_table(tbl)
    
    if col_widths:
        for row in tbl.rows:
            for c_idx, w in enumerate(col_widths):
                row.cells[c_idx].width = w
                
    ch4_element.addprevious(tbl._tbl)
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(12)
    ch4_element.addprevious(p_spacer._p)
    return tbl

print("Inserting clean, structured Chapter 3 content with high-res diagram images...")

# 3.1 Class Diagram
add_heading2("3.1} Class Diagram.")
add_p("The system's domain model is organized according to object-oriented software engineering principles. Below is the structural UML Class Diagram defining the principal entities, their private attributes, member methods, and architectural associations:")
add_diagram("3.1_Class_Diagram.png", "Figure 3.1: Comprehensive UML Domain Class Diagram", width_in=6.0)

# 3.2 Use-Case Diagram
add_heading2("3.2} Use-Case Diagram.")
add_p("The Use-Case model delineates system interactions between external actors (Student and Educator / Admin) and the core functional capabilities of the QuizNova-AI platform:")
add_diagram("3.2_Use_Case_Diagram.png", "Figure 3.2: System Use-Case Diagram for Admin / Educator and Student Roles", width_in=6.0)

# 3.2.1 Admin Use-Case
add_heading3("3.2.1} Admin Use-Case.")
add_p("The Admin / Educator Use-Case workflow specifically governs administrative supervision, examination moderation, and curriculum control across the platform:")
add_bullet("Curriculum Administration: Create, organize, and manage assessment categories matching Maharashtra State Board syllabi for SSC (Classes 9-10) and HSC (Classes 11-12 Science, Commerce, Arts).", "1. ")
add_bullet("Assessment Authoring: Compose dynamic question items, specify options, define answer keys, allocate marks, and tag questions with deterministic Bloom's taxonomy cognitive levels.", "2. ")
add_bullet("Examination Enforcement Configuration: Define exam time limits, passing scores, allowed tab-switch violation limits, and fullscreen mandates.", "3. ")
add_bullet("Evaluation & Audit Verification: Access cohort grade distributions, analyze question difficulty ratings, and audit student security violation logs.", "4. ")
add_diagram("3.2.1_Admin_Use_Case.png", "Figure 3.3: Admin Dedicated Use-Case & Assessment Creation Flow", width_in=6.0)

# 3.3 Sequence Diagram
add_heading2("3.3} Sequence Diagram.")
add_p("The following Sequence Diagram illustrates the high-concurrency interaction among the Candidate Browser, Anti-Cheating Proctor Hook, Express API Controller, Prisma ORM, and Single Unified Database during an assessment lifecycle:")
add_diagram("3.3_Sequence_Diagram.png", "Figure 3.4: High-Concurrency Sequence Diagram for Proctored Assessment Lifecycle", width_in=6.2)

# 3.4 Activity Diagram
add_heading2("3.4} Activity Diagram.")
add_p("The Activity Diagram models the dynamic operational flow from user authentication through proctored exam delivery, integrity enforcement, automated grading, and cognitive scoring:")
add_diagram("3.4_Activity_Diagram.png", "Figure 3.5: System Activity Diagram for Exam Execution and Anti-Cheating Supervision", width_in=5.8)

# 3.5 Data Flow Diagrams
add_heading2("3.5} Data flow Diagram (DFD).")
add_p("Data Flow Diagrams illustrate how information transitions through the system's operational processes, external actors, and unified database stores across various levels of abstraction.")

# 3.5.1 Level 0 DFD
add_heading3("3.5.1} Level 0 DFD (Context Level Diagram)")
add_p("The Context Diagram defines the global external entities interacting with the QuizNova-AI system boundary:")
add_diagram("3.5.1_Level_0_DFD.png", "Figure 3.6: Data Flow Diagram (DFD) Level 0 – Context Level System Boundary", width_in=5.8)

# 3.5.2 Level 1 DFD
add_heading3("3.5.2} Level 1 DFD (Functional Decomposition)")
add_p("Decomposes the system into its primary subsystems: Process 1.0 (Authentication & User Management), Process 2.0 (Assessment Authoring & Curriculum Management), Process 3.0 (Proctored Exam Execution & Telemetry Supervision), Process 4.0 (Diagnostic Scoring & Analytics Reporting), and Process 5.0 (Digital Textbook & 36-Year PYQ Library Delivery).")
add_diagram("3.5.2_Level_1_DFD.png", "Figure 3.7: Data Flow Diagram (DFD) Level 1 – Subsystem Functional Decomposition", width_in=6.0)

# 3.5.3 Level 2A DFD
add_heading3("3.5.3} Level 2A DFD (Assessment Creation & Bloom Engine)")
add_p("Focuses on how educators author assessments, tag cognitive levels, and commit records atomically to the database:")
add_bullet("Process 2.1: Validate examination parameters (Title, Duration, Standard, Marks).")
add_bullet("Process 2.2: Parse question items and options JSON schema.")
add_bullet("Process 2.3: Associate Bloom's Taxonomy cognitive tier with each question.")
add_bullet("Process 2.4: Atomically commit records into SQLite via Prisma relational transaction.")
add_diagram("3.5.3_Level_2A_DFD.png", "Figure 3.8: Data Flow Diagram (DFD) Level 2A – Assessment Creation & Bloom Engine", width_in=5.8)

# 3.5.4 Level 2B DFD
add_heading3("3.5.4} Level 2B DFD (Exam Execution & Proctoring Sub-Process)")
add_p("Focuses on runtime assessment delivery, anti-cheat detection, and exam telemetry enforcement:")
add_bullet("Process 3.1: Verify candidate eligibility and lock fullscreen execution.")
add_bullet("Process 3.2: Execute hardware clock-aligned countdown timer.")
add_bullet("Process 3.3: Intercept Page Visibility API events and document blur triggers.")
add_bullet("Process 3.4: Accumulate violation telemetry; execute auto-submit upon reaching limit.")
add_bullet("Process 3.5: Transmit answers payload to evaluation engine.")
add_diagram("3.5.4_Level_2B_DFD.png", "Figure 3.9: Data Flow Diagram (DFD) Level 2B – Exam Execution & Proctoring Sub-Process", width_in=5.8)

# 3.5.6 Gantt Chart
add_heading3("3.5.6} Gantt Chart (Project SDLC Schedule)")
add_p("The project schedule followed an Agile timeline spanning five engineering sprints across 16 weeks of structured development, testing, and deployment:")
add_diagram("3.5.6_Gantt_Chart.png", "Figure 3.10: Project SDLC Gantt Chart & Engineering Sprint Schedule", width_in=6.2)

# Gantt Schedule Table
gantt_headers = ["Task / Sprint ID", "Sprint Description", "Duration", "Target Milestone", "Status"]
gantt_rows = [
    ["Sprint 1", "System Architecture, Database Modeling & Prisma Schema", "Weeks 1-4", "Unified SQLite DB & Types", "Completed"],
    ["Sprint 2", "1-Click Google OAuth & Direct Academic Authentication", "Weeks 5-8", "Zero-OTP Secure Auth Hub", "Completed"],
    ["Sprint 3", "Exam Player, Fullscreen Proctor & Offline IndexedDB Cache", "Weeks 9-12", "Anti-Cheat & Offline Hub", "Completed"],
    ["Sprint 4", "41 Balbharati Textbooks & 36-Year Subjective PYQ Archive", "Weeks 13-14", "Digital Library & State PYQs", "Completed"],
    ["Sprint 5", "Comprehensive QA, Unit Testing, Vercel CI/CD & Documentation", "Weeks 15-16", "Production Deployment", "Completed"]
]
add_tbl(gantt_headers, gantt_rows, [Inches(1.1), Inches(2.5), Inches(0.9), Inches(1.5), Inches(0.8)])

# 3.6 Data Modelling
add_heading2("3.6} Data Modelling.")

add_heading3("3.6.1} Main Database Entities")
add_p("The relational schema is composed of four foundational entities operating within a single unified database:")
add_bullet("Users Entity: Stores candidate and educator identity records, encrypted credentials, and roles.")
add_bullet("Quizzes Entity: Maintains examination definitions, durations, and passing standards.")
add_bullet("Questions Entity: Contains individual assessment items, options JSON, answer keys, and cognitive Bloom tiers.")
add_bullet("Quiz Attempts Entity: Records student attempt metrics, total scores, violation counts, and Bloom diagnostics.")

add_heading3("3.6.2} Purpose of Data Model")
add_p("The data model enforces relational normalization (Third Normal Form - 3NF) to eliminate data redundancy, maintain referential integrity with cascading deletes, ensure atomic transactions across high concurrency, and provide sub-10ms query performance on SQLite.")

# 3.6.3 Users Table
add_heading3("3.6.3} Data Dictionary – Users Table (users)")
users_headers = ["Column Name", "Data Type", "Constraint", "Description"]
users_rows = [
    ["id", "VARCHAR(36)", "PRIMARY KEY", "Unique UUID identifier for the user account"],
    ["name", "VARCHAR(100)", "NOT NULL", "Full name of the student or educator"],
    ["email", "VARCHAR(150)", "UNIQUE, NOT NULL", "Primary email address for login"],
    ["passwordHash", "VARCHAR(255)", "NOT NULL", "Salted bcrypt hash of user password"],
    ["role", "VARCHAR(20)", "NOT NULL", "User authorization role (STUDENT, EDUCATOR)"],
    ["phoneNumber", "VARCHAR(20)", "NULLABLE", "Optional contact phone number"],
    ["createdAt", "DATETIME", "DEFAULT CURRENT_TIMESTAMP", "Account creation timestamp"]
]
add_tbl(users_headers, users_rows, [Inches(1.4), Inches(1.3), Inches(1.5), Inches(2.3)])

# 3.6.4 Quizzes Table
add_heading3("3.6.4} Data Dictionary – Quizzes Table (quizzes)")
quizzes_headers = ["Column Name", "Data Type", "Constraint", "Description"]
quizzes_rows = [
    ["id", "VARCHAR(36)", "PRIMARY KEY", "Unique UUID identifier for the assessment"],
    ["title", "VARCHAR(200)", "NOT NULL", "Title of the assessment"],
    ["description", "TEXT", "NULLABLE", "Detailed description and student instructions"],
    ["educatorId", "VARCHAR(36)", "FOREIGN KEY (Users.id)", "UUID of the authoring educator"],
    ["timeLimitMinutes", "INTEGER", "NOT NULL, DEFAULT 30", "Allocated exam time limit in minutes"],
    ["passingScore", "FLOAT", "NOT NULL, DEFAULT 40.0", "Minimum percentage required to pass"],
    ["isPublished", "BOOLEAN", "DEFAULT TRUE", "Visibility flag for student enrollment"]
]
add_tbl(quizzes_headers, quizzes_rows, [Inches(1.5), Inches(1.3), Inches(1.5), Inches(2.2)])

# 3.6.5 Quiz Attempts Table
add_heading3("3.6.5} Data Dictionary – Quiz Attempts Table (quiz_attempts)")
attempts_headers = ["Column Name", "Data Type", "Constraint", "Description"]
attempts_rows = [
    ["id", "VARCHAR(36)", "PRIMARY KEY", "Unique UUID identifier for the attempt record"],
    ["quizId", "VARCHAR(36)", "FOREIGN KEY (Quizzes.id)", "UUID of the attempted assessment"],
    ["studentId", "VARCHAR(36)", "FOREIGN KEY (Users.id)", "UUID of the participating candidate"],
    ["scoreObtained", "FLOAT", "NOT NULL", "Total marks scored by the student"],
    ["totalMarks", "FLOAT", "NOT NULL", "Maximum possible marks for the assessment"],
    ["percentage", "FLOAT", "NOT NULL", "Normalized percentage score"],
    ["bloomData", "JSON / TEXT", "NOT NULL", "6-tier Bloom's taxonomy performance breakdown"],
    ["violationCount", "INTEGER", "DEFAULT 0", "Count of detected tab switches and blur events"],
    ["completedAt", "DATETIME", "DEFAULT CURRENT_TIMESTAMP", "Submission timestamp"]
]
add_tbl(attempts_headers, attempts_rows, [Inches(1.4), Inches(1.3), Inches(1.6), Inches(2.2)])

# 3.6.6 Relationships Between Tables
add_heading3("3.6.6} Relationship between tables")
add_bullet("Users to Quizzes: 1-to-Many relationship (Users.id -> Quizzes.educatorId). One educator authors multiple quizzes.")
add_bullet("Quizzes to Questions: 1-to-Many relationship with cascading deletes (Quizzes.id -> Questions.quizId). One quiz contains multiple questions.")
add_bullet("Quizzes to QuizAttempts: 1-to-Many relationship (Quizzes.id -> QuizAttempts.quizId). One quiz records attempts from multiple students.")
add_bullet("Users to QuizAttempts: 1-to-Many relationship (Users.id -> QuizAttempts.studentId). One student submits multiple quiz attempts.")

# 3.7 E-R Diagram
add_heading2("3.7} E-R Diagram.")
add_p("Below is the comprehensive Entity-Relationship Diagram representing the unified database schema, relational cardinalities, attributes, and primary/foreign key mappings:")
add_diagram("3.7_ER_Diagram.png", "Figure 3.11: Entity-Relationship (E-R) Diagram (Single Unified Relational Schema)", width_in=6.0)

print(f"Saving final formatted Word document to: {output_project_docx}")
doc.save(output_project_docx)

print(f"Saving copy directly to user Desktop: {output_desktop_docx}")
doc.save(output_desktop_docx)

print("SUCCESS: Both documents generated successfully!")
